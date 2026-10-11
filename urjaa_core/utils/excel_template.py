"""
Generates the unified bulk-upload Excel template for products.

Sheet layout:
  1. products  – one row per product (human-readable names; slug auto-generated)
  2. variants  – one row per variant linked by product_name
  3. stones    – optional, one row per stone line linked by product_name
  4. instructions – field reference plus the current valid categories,
     subcategories, metal combinations, genders and stone types, read from the DB at
     download time so staff always see what's actually in the catalog (K-01).
"""

import io
import openpyxl
from fastapi import HTTPException
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from urjaa_core.repositories.admin.admin_management_repository import AdminManagementRepository


# K-01: row 2 of the products/variants/stones sheets is always the filled-in example
# -- real data starts here, regardless of what row 2 contains. Simpler and
# more predictable than trying to detect "is this still the example row".
EXAMPLE_ROW_NUMBER = 2

_GOLD = "C9A14A"
_DARK_BG = "111111"
_ROW_BG = "1A1A1A"
_HEADER_FILL = PatternFill("solid", fgColor=_GOLD)
_ROW_FILL = PatternFill("solid", fgColor=_ROW_BG)
_HEADER_FONT = Font(bold=True, color=_DARK_BG, size=10)
_EXAMPLE_FONT = Font(italic=True, color="888888", size=9)
_LABEL_FONT = Font(bold=True, color=_GOLD, size=10)
_BODY_FONT = Font(color="CCCCCC", size=9)
_REQUIRED_FONT = Font(bold=True, color="FF6B6B", size=9)
_THIN = Side(style="thin", color="333333")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_cell(cell, value: str):
    cell.value = value
    cell.fill = _HEADER_FILL
    cell.font = _HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    cell.border = _BORDER


def _example_cell(cell, value: str):
    cell.value = value
    cell.font = _EXAMPLE_FONT
    cell.fill = _ROW_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    cell.border = _BORDER


def _set_col_widths(ws, widths: list[int]):
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


# ---------------------------------------------------------------------------
# Unified column definitions
# ---------------------------------------------------------------------------

# Products sheet — one row per unique product
# (*) = required; everything else is optional
PRODUCT_COLUMNS = [
    # (column_name, is_required, example_value, description)
    ("product_name",    True,  "Gold Filigree Necklace",         "Unique display name; slug is auto-generated"),
    ("category_name",   True,  "Necklaces",                       "Exact category name as shown in your catalog (case-insensitive)"),
    ("subcategory_name",False, "Choker Necklaces",                "Subcategory name under the given category (optional)"),
    ("gender",          False, "unisex",                          "men | women | unisex | kids (case-insensitive; default: unisex)"),
    ("description",     False, "Handcrafted 22K gold necklace.",  "Short plain-text description"),
    ("status",          False, "active",                          "active | draft | hidden | archived — leave blank = draft (inactive, not shown on the storefront)"),
    ("featured",        False, "false",                           "true | false  (default: false)"),
    ("customizable",    False, "false",                           "true | false  (default: false)"),
]

# Variants sheet — one or more rows per product (linked by product_name)
VARIANT_COLUMNS = [
    ("product_name",       True,  "Gold Filigree Necklace",  "Must exactly match a product_name in the Products sheet"),
    ("variant_sku",        False, "GFN-22K-16IN",            "Leave blank to auto-generate a SKU"),
    ("metal",              False, "22K Yellow Gold",         "Exact metal combination name from your catalog (case-insensitive) — see VALID METAL COMBINATIONS below"),
    ("size_value",         False, "6",                       "Size/length value, in the unit the subcategory uses — see VALID SUBCATEGORIES below. Leave blank if the subcategory has no size"),
    ("spec_note",          False, "Engraved initials",       "One-off display note, max 200 characters"),
    ("metal_weight_grams", False, "8.5",                     "Decimal grams of metal (used in computed price)"),
    ("making_charges",     False, "1500",                    "Fixed making charge in INR"),
    ("price_override",     False, "5500",                    "Fixed price in INR; overrides metal-rate computation if set"),
    ("stock_quantity",     False, "10",                      "Integer >= 0  (default: 0)"),
]


# Stones sheet — optional, one row per stone line (linked by product_name).
# Stones belong to the product (product_stones has no variant link), so every
# variant of that product carries them -- same as the product editor.
STONE_COLUMNS = [
    ("product_name",         True,  "Gold Filigree Necklace", "Must exactly match a product_name in the Products sheet. Several stones = several rows"),
    ("stone_type",           True,  "Diamond",                "Stone name from your Stones master (case-insensitive) — see VALID STONE TYPES below"),
    ("quantity",             False, "12",                     "Number of stones in this line (whole number)"),
    ("total_carat_weight",   False, "0.36",                   "Total carat weight of this line (decimal)"),
    ("cost",                 False, "18000",                  "Total cost of this line in INR (quantity x unit price) — added to the computed price"),
    ("cut",                  False, "Excellent",              "Cut / shape description, max 50 characters"),
    ("clarity",              False, "VS1",                    "Clarity grade, max 20 characters"),
    ("color",                False, "F",                      "Colour grade or colour, max 20 characters"),
    ("origin",               False, "Natural",                "Origin, max 100 characters"),
    ("certificate_number",   False, "IGI-123456789",          "Certificate number, max 100 characters"),
    ("certification_agency", False, "IGI",                    "GIA | IGI | HRD | BIS"),
]


def _build_data_sheet(wb: openpyxl.Workbook, name: str, columns: list, widths: list[int]):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20

    # Row 1: column headers (hover a header for its description); row 2: example values
    for col_idx, (col_name, required, example, desc) in enumerate(columns, start=1):
        header = ws.cell(row=1, column=col_idx)
        _header_cell(header, f"{'* ' if required else ''}{col_name}")
        header.comment = Comment(desc, "Urjaa")
        _example_cell(ws.cell(row=2, column=col_idx), example)

    _set_col_widths(ws, widths)
    return ws


def _build_instructions_sheet(wb: openpyxl.Workbook, db: Session):
    ws = wb.create_sheet("instructions")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 32

    def _title(row, text):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, color=_GOLD, size=12)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 26

    def _section(row, text):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="222222")
        cell.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 22

    def _field_row(row, col_name, required, example, desc):
        ws.cell(row=row, column=1, value=col_name).font = Font(
            bold=required, color="FFCCAA" if required else "CCCCCC", size=9
        )
        ws.cell(row=row, column=2, value="REQUIRED" if required else "optional").font = Font(
            color="FF6B6B" if required else "666666", size=8, italic=not required
        )
        ws.cell(row=row, column=3, value=desc).font = _BODY_FONT
        ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=4, value=f"e.g.  {example}" if example else "").font = _EXAMPLE_FONT
        ws.row_dimensions[row].height = 18

    def _list_row(row, value, note=""):
        ws.cell(row=row, column=1, value=value).font = Font(color="C9A14A", size=9)
        if note:
            ws.cell(row=row, column=3, value=note).font = _BODY_FONT
        ws.row_dimensions[row].height = 16

    r = 1
    _title(r, "BULK UPLOAD — FIELD REFERENCE")
    r += 1
    ws.cell(row=r, column=1, value="* = required column   |   All other columns are optional").font = Font(
        color="888888", size=9, italic=True
    )
    r += 2

    # General rules
    _section(r, "HOW TO USE THIS TEMPLATE")
    r += 1
    rules = [
        "Fill the 'products' sheet — one row per unique product.",
        "Fill the 'variants' sheet — one or more rows per product, linked by product_name.",
        "Optional: fill the 'stones' sheet — one row per stone line (e.g. 12 diamonds + 1 ruby = 2 rows), "
        "linked by product_name. Stones apply to every variant of that product. Leave it empty for plain-metal pieces.",
        "status: leave blank = draft (inactive, NOT shown on the storefront). Write active to publish.",
        f"Upload this .xlsx file directly — no need to export to CSV (CSV uploads still work too, "
        f"if you build your own one-sheet file).",
        f"Row {EXAMPLE_ROW_NUMBER} of every sheet is a filled-in example. Uploading the .xlsx directly "
        f"always ignores row {EXAMPLE_ROW_NUMBER} — enter your data starting at row {EXAMPLE_ROW_NUMBER + 1}.",
        "category_name must match a name in your catalog exactly (case-insensitive). No slug needed.",
        "subcategory_name is optional. If given, it must belong to the specified category.",
        "gender is optional — men, women, unisex or kids (case-insensitive). Leave blank for unisex.",
        "Product slug is auto-generated from product_name — never provide a slug column.",
        "variant_sku is optional — leave blank to auto-generate one; must be globally unique if given.",
        "metal must match a metal combination's display name exactly (case-insensitive), e.g. \"22K Yellow Gold\" — see VALID METAL COMBINATIONS below.",
        "size_value only applies to subcategories with a size label configured — see VALID SUBCATEGORIES below.",
        "Rows with the same product_name create variants of the same product.",
        "Use price_override to set a fixed price; leave blank to compute via metal rate.",
    ]
    for rule in rules:
        ws.cell(row=r, column=1, value=f"  {rule}").font = Font(color="AAAAAA", size=9)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 16
        r += 1

    r += 1
    _section(r, "PRODUCTS SHEET FIELDS")
    r += 1
    for col_name, required, example, desc in PRODUCT_COLUMNS:
        _field_row(r, col_name, required, example, desc)
        r += 1

    r += 1
    _section(r, "VARIANTS SHEET FIELDS")
    r += 1
    for col_name, required, example, desc in VARIANT_COLUMNS:
        _field_row(r, col_name, required, example, desc)
        r += 1

    r += 1
    _section(r, "STONES SHEET FIELDS  (optional sheet)")
    r += 1
    for col_name, required, example, desc in STONE_COLUMNS:
        _field_row(r, col_name, required, example, desc)
        r += 1

    r += 1
    _section(r, "STATUS VALUES")
    r += 1
    for val, desc in [
        ("draft", "Default when status is left blank — inactive, not visible on the storefront"),
        ("active", "Product is live and visible to customers"),
        ("hidden", "Inactive — kept out of the storefront"),
        ("archived", "Soft-deleted — hidden from all views"),
    ]:
        ws.cell(row=r, column=1, value=val).font = Font(color="C9A14A", bold=True, size=9)
        ws.cell(row=r, column=2, value=desc).font = _BODY_FONT
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 16
        r += 1

    # K-01: the current catalog, read live so staff never fill in a name that
    # no longer exists.
    r += 1
    _section(r, "VALID GENDERS")
    r += 1
    for gender in AdminManagementRepository.get_genders(db):
        _list_row(r, gender.name)
        r += 1

    r += 1
    _section(r, "VALID CATEGORIES")
    r += 1
    categories = AdminManagementRepository.get_categories(db)
    for category in categories:
        _list_row(r, category.name)
        r += 1
    if not categories:
        _list_row(r, "(none configured yet)")
        r += 1

    r += 1
    _section(r, "VALID SUBCATEGORIES  (category — size label / unit, if any)")
    r += 1
    subcategories = AdminManagementRepository.get_subcategories(db)
    for subcategory in subcategories:
        category_name = subcategory.category.name if subcategory.category else "?"
        size_info = (
            f"{subcategory.size_label} / {subcategory.size_unit}" if subcategory.size_label and subcategory.size_unit
            else (subcategory.size_label or "no sizing")
        )
        _list_row(r, f"{subcategory.name}  —  {category_name}", size_info)
        r += 1
    if not subcategories:
        _list_row(r, "(none configured yet)")
        r += 1

    r += 1
    _section(r, "VALID METAL COMBINATIONS")
    r += 1
    metal_names = sorted({metal.display_name for metal in AdminManagementRepository.get_metals(db) if metal.display_name})
    for name in metal_names:
        _list_row(r, name)
        r += 1
    if not metal_names:
        _list_row(r, "(none configured yet — generate combinations on the Metals admin page)")
        r += 1

    r += 1
    _section(r, "VALID STONE TYPES")
    r += 1
    stones = AdminManagementRepository.get_stones(db)
    for stone in stones:
        _list_row(r, stone.name)
        r += 1
    if not stones:
        _list_row(r, "(none configured yet — add them on the Stones admin page)")
        r += 1


def generate_product_template(db: Session) -> bytes:
    """Return the unified product bulk-upload Excel template as raw bytes.

    Needs a DB session (K-01): the instructions sheet lists the catalog's
    current valid categories, subcategories, metal combinations, genders and
    stone types as they exist at download time.
    """
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    _build_data_sheet(wb, "products", PRODUCT_COLUMNS, [30, 26, 26, 14, 44, 10, 10, 14])
    _build_data_sheet(wb, "variants", VARIANT_COLUMNS, [30, 18, 20, 14, 24, 20, 18, 16, 14])
    _build_data_sheet(wb, "stones", STONE_COLUMNS, [30, 18, 10, 18, 12, 14, 10, 10, 14, 20, 20])
    _build_instructions_sheet(wb, db)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def parse_bulk_upload_workbook(
    content: bytes,
) -> tuple[list[tuple[int, dict[str, str]]], ...]:
    """K-01: read an uploaded bulk-upload .xlsx (the file generate_product_template
    produced, filled in) into raw (row_number, {header: value}) pairs for the
    'products', 'variants' and (optional -- older templates don't have it)
    'stones' sheets.

    Row EXAMPLE_ROW_NUMBER (2) of each sheet is always skipped, regardless of
    content -- it's the template's filled-in example; see the instructions
    sheet. Values come back as stripped strings (same shape a CSV row
    produces) so the caller can feed them straight into the existing
    product_name-keyed joining and validation.

    Raises HTTPException(400) for a file that isn't structurally the
    template (unreadable, or missing a sheet) -- that's a different file,
    not a data problem the normal per-row errors can describe.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=400, detail="Could not read this .xlsx file — is it the downloaded template?"
        ) from exc

    try:
        missing = [name for name in ("products", "variants") if name not in wb.sheetnames]
        if missing:
            raise HTTPException(
                status_code=400,
                detail=f"Missing sheet(s): {', '.join(missing)} — download the current template and fill it in.",
            )

        def _sheet_rows(sheet_name: str) -> list[tuple[int, dict[str, str]]]:
            ws = wb[sheet_name]
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            # _header_cell prefixes a required column with "* " -- strip it
            # back off so the key matches the column name (e.g. "product_name").
            headers = [str(cell).strip().lstrip("* ").strip() if cell is not None else "" for cell in header_row]

            rows: list[tuple[int, dict[str, str]]] = []
            data_rows = ws.iter_rows(min_row=EXAMPLE_ROW_NUMBER + 1, values_only=True)
            for row_number, row in enumerate(data_rows, start=EXAMPLE_ROW_NUMBER + 1):
                if all(cell is None or str(cell).strip() == "" for cell in row):
                    continue
                values = {
                    headers[i]: ("" if cell is None else str(cell).strip())
                    for i, cell in enumerate(row)
                    if i < len(headers) and headers[i]
                }
                rows.append((row_number, values))
            return rows

        stones = _sheet_rows("stones") if "stones" in wb.sheetnames else []
        return _sheet_rows("products"), _sheet_rows("variants"), stones
    finally:
        wb.close()
