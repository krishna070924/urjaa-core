"""K-01 self-check against the live dev database, as the admin app role.
The service commits internally, so (same reasoning as
tests/test_sku_autogeneration.py) the session is joined to an outer
transaction via savepoints and everything is rolled back at the end.

Covers:
1. A row with a blank variant_sku, a metal name and a gender name ->
   generated SKU (H-04), metal_id resolved from the Metal row (H-10 rule),
   gender_id resolved by name (H-11 rule). Blank gender on another row ->
   the same default create_product uses.
2. Unknown metal name -> row error naming the value and listing valid metals.
3. Unknown gender name -> row error naming the value and listing valid genders.
4. size_value on a subcategory with no size_label -> row error saying so.
5. A file still using the old base_metal_id/stone_cost columns -> rejected
   up front with a clear message (the chosen old-file behaviour; see K-01
   report for why reject beats a per-row warning).
6. generate_product_template(db) returns bytes whose instructions sheet
   lists current categories, metals and genders from the DB.

Run: .venv/bin/python tests/test_excel_import_v2.py
"""
import os
import uuid

os.environ.setdefault("DATABASE_URL", "postgresql://postgres:postgres@localhost:5432/urjaa")

import openpyxl
import io
from sqlalchemy import text
from sqlalchemy.orm import Session

from urjaa_core.core.database import engine
from urjaa_core.models import Product, ProductVariant, Store
from urjaa_core.services.admin.admin_management_service import AdminManagementService as svc
from urjaa_core.utils.excel_template import generate_product_template


def main() -> None:
    conn = engine.connect()
    outer = conn.begin()
    db = Session(bind=conn, join_transaction_mode="create_savepoint")
    try:
        db.execute(text("SET ROLE urjaa_admin_svc"))
        store = db.query(Store).first()
        assert store is not None, "dev DB has no seeded store"
        suffix = uuid.uuid4().hex[:8]

        # --- 1. Happy path: blank SKU, metal by name, gender by name ---
        name_a = f"K01 Ring A {suffix}"
        result = svc.bulk_upload_products_from_csv_rows(
            db, store.id,
            [{
                "product_name": name_a,
                "category_name": "Rings",
                "subcategory_name": "Cocktail Rings",
                "gender": "Men",
                "metal": "22k yellow gold",  # case/space-insensitive match
                "size_value": "6",
                "variant_sku": "",
                "stock_quantity": "3",
            }],
        )
        assert result["success"] and result["inserted_products"] == 1 and result["inserted_variants"] == 1, result

        product = db.query(Product).filter(Product.name == name_a).one()
        variant = db.query(ProductVariant).filter(ProductVariant.product_id == product.id).one()
        assert variant.sku_code and variant.sku_code.startswith("RIN-"), variant.sku_code
        assert variant.metal_id is not None and variant.metal.display_name == "22K Yellow Gold", variant.metal_id
        assert variant.base_metal_id == variant.metal.base_metal_id, "legacy base_metal_id not set from Metal row"
        assert product.gender_id is not None and product.gender.name == "Men", product.gender_id
        assert variant.size_value == "6", variant.size_value
        print(f"blank SKU + metal/gender by name  -> OK  sku={variant.sku_code} metal_id={variant.metal_id} gender={product.gender.name}")

        # Blank gender column -> same default as create_product (unisex).
        name_b = f"K01 Ring B {suffix}"
        result_b = svc.bulk_upload_products_from_csv_rows(
            db, store.id,
            [{"product_name": name_b, "category_name": "Rings", "subcategory_name": "Cocktail Rings"}],
        )
        assert result_b["success"], result_b
        product_b = db.query(Product).filter(Product.name == name_b).one()
        assert product_b.gender.name == "Unisex", product_b.gender.name
        print("blank gender column               -> OK  defaults to Unisex")

        # --- 2. Unknown metal -> row error naming it + listing valid ones ---
        result = svc.bulk_upload_products_from_csv_rows(
            db, store.id,
            [{
                "product_name": f"K01 Ring C {suffix}",
                "category_name": "Rings",
                "subcategory_name": "Cocktail Rings",
                "metal": "22K Rose",
            }],
        )
        assert not result["success"], result
        msg = result["errors"][0]["errors"][0]
        assert "22K Rose" in msg and "Valid:" in msg and "22K Yellow Gold" in msg, msg
        print(f"unknown metal                     -> OK  {msg[:70]}...")

        # --- 3. Unknown gender -> row error naming it + listing valid ones ---
        result = svc.bulk_upload_products_from_csv_rows(
            db, store.id,
            [{
                "product_name": f"K01 Ring D {suffix}",
                "category_name": "Rings",
                "subcategory_name": "Cocktail Rings",
                "gender": "nonbinary",
            }],
        )
        assert not result["success"], result
        msg = result["errors"][0]["errors"][0]
        assert "nonbinary" in msg and "Valid:" in msg and "Unisex" in msg, msg
        print(f"unknown gender                    -> OK  {msg[:70]}...")

        # --- 4. size_value on a subcategory with no size_label -> error ---
        result = svc.bulk_upload_products_from_csv_rows(
            db, store.id,
            [{
                "product_name": f"K01 Earring A {suffix}",
                "category_name": "Earrings",
                "subcategory_name": "Hoops",
                "size_value": "5",
            }],
        )
        assert not result["success"], result
        msg = result["errors"][0]["errors"][0]
        assert "Hoops" in msg and "sizes" in msg, msg
        print(f"size on a size-less subcategory   -> OK  {msg}")

        # --- 5. Old template columns -> rejected with a clear message ---
        result = svc.bulk_upload_products_from_csv_rows(
            db, store.id,
            [{
                "product_name": f"K01 Ring E {suffix}",
                "category_name": "Rings",
                "base_metal_id": "1",
                "stone_cost": "500",
            }],
        )
        assert not result["success"] and result["inserted_products"] == 0, result
        msg = result["errors"][0]["errors"][0]
        assert "base_metal_id" in msg and "stone_cost" in msg and "new template" in msg.lower(), msg
        print(f"old template columns              -> OK  rejected up front")

        # --- 6. Template generation lists the live catalog ---
        data = generate_product_template(db)
        assert isinstance(data, bytes) and len(data) > 1000, len(data)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["products", "variants", "instructions"], wb.sheetnames
        all_values = {str(cell.value) for row in wb["instructions"].iter_rows() for cell in row if cell.value is not None}
        assert "22K Yellow Gold" in all_values, "metal combinations missing from instructions sheet"
        assert "Rings" in all_values, "categories missing from instructions sheet"
        assert "Unisex" in all_values, "genders missing from instructions sheet"
        assert any("Cocktail Rings" in v for v in all_values), "subcategories missing from instructions sheet"
        print(f"template generation                -> OK  {len(data)} bytes, instructions lists current catalog")

        print("ALL OK")
    finally:
        db.close()
        outer.rollback()
        conn.close()


if __name__ == "__main__":
    main()
